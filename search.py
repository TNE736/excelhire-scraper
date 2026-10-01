"""Excel Hire -> Job Seekers -> AI Search -> External: paste a JD, run the search, save the top
candidates to Excel.

No LLM and no credits. It only reads what the results table already shows. It never clicks
Reveal (the one thing that spends credits), and it reads the credit counter before and after
each run so a change would be reported, not silent.

Setup, once (PowerShell, in this folder):
    pip install -r requirements.txt
    python -m playwright install chromium
    copy .env.example .env      then put your Excel Hire email and password in .env

Run:
    python search.py jd.txt                      top 20, Best Fit
    python search.py jd.txt --count 25 --mode Exact
    python search.py jd.txt --show               show the browser (first login, or to watch it)
    python search.py jd.txt --csv out.csv        also write the LeadOps CSV (see CSV_COLUMNS)

The login is kept in browser-profile/ between runs. The Excel file goes to output/.

For programs that run it (LeadOps): each step prints a `PROGRESS <step>` line as it starts, and
the exit code says how it ended (see EXIT_CODES).
"""

from __future__ import annotations

import argparse
import csv
import os
import re
import sys
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from playwright.sync_api import Locator, Page, sync_playwright
from playwright.sync_api import TimeoutError as PlaywrightTimeout

HERE = Path(__file__).resolve().parent

# How a run ends, for programs that start it. Printed messages stay the same for people.
EXIT_CODES = {
    "ok": 0,
    "bad_input": 1,  # missing JD file, JD too short
    "login_failed": 2,
    "table_changed": 3,  # Excel Hire changed its results table; nothing was saved
    "credits_changed": 4,  # files were saved, but the credit counter moved during the run
    "timed_out": 5,  # Excel Hire did not answer in time
}

# The LeadOps CSV (--csv): one row per candidate, columns in this order. No email or phone:
# they sit behind Reveal, which this scraper never clicks.
CSV_COLUMNS = (
    "first_name",
    "last_name",
    "linkedin_url",
    "location",
    "company",
    "skills",
    "experience",
    "education",
    "match_score",
    "match_reasons",
)
PROFILE_FOLDER = HERE / "browser-profile"  # keeps the login between runs
OUTPUT_FOLDER = HERE / "output"

SEARCH_MODES = ("Exact", "Best Fit", "Skill Match")
MIN_JD_LENGTH = 10  # Excel Hire's own minimum for the query box
ROWS_ON_FIRST_PAGE = 25  # what one External search shows before paging (seen 24 Sept 2026)

# Seconds. Parsing the JD and ranking the search both happen on Excel Hire's side.
PAGE_TIMEOUT = 60
FILTERS_TIMEOUT = 90
RESULTS_TIMEOUT = 120
LOGIN_TIMEOUT = 180

# The results table, as Excel Hire builds it. Each cell names itself by class, so a column
# moving around does not break anything; a class disappearing makes the run stop, not guess.
RESULT_ROW = "table.c-table tbody tr"
CELL = {
    "name": "td.name-cell",
    "job_title": "td.job-title-cell",
    "company": "td.company-cell",
    "experience": "td.experience-cell",
    "skills": "td.skills-cell",
    "education": "td.education-cell",
    "score": "td.score-cell",
}
EXPAND_LINK = re.compile(r"^\s*\+\d+ more\s*$")  # "+8 more": expands in the page, costs nothing
EMPTY_CELL = "-"

# The login page (seen 24 Sept 2026): one "name@company.com" box and a Login button first; the
# password box comes after.
EMAIL_INPUT = (
    "input[type=email], input[placeholder*='@'], input[formcontrolname*=email i], "
    "input[name*=email i]"
)
LOGIN_BUTTON = "button[type=submit], button:has-text('Login'), button:has-text('Sign in')"

# The Candidates sheet: (header, width, source field). The most useful columns come first and
# stay in view (frozen) when scrolling right.
EXCEL_COLUMNS = (
    ("#", 5, None),
    ("Name", 24, "Name"),
    ("Score", 8, "Score (%)"),
    ("Job Title", 30, "Job Title"),
    ("Company", 24, "Company"),
    ("Location", 24, "Location"),
    ("Experience", 55, "Experience"),
    ("Skills", 50, "Skills"),
    ("Education", 42, "Education"),
    ("LinkedIn", 14, "LinkedIn"),
    ("Why matched", 60, "Why matched"),
)
FROZEN_UP_TO = "D2"  # keeps the header row and #, Name, Score on screen
ROW_HEIGHT = 96  # points, about 7 lines; a longer cell keeps all its text - widen the row
FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="4C1D95")  # Excel Hire's purple, darkened for contrast
BAND_FILL = PatternFill("solid", fgColor="F5F3FF")  # every other row, so rows are easy to follow
LINK_COLOR = "1D4ED8"
TOP_LEFT = Alignment(wrap_text=True, vertical="top")


class RunStopped(Exception):
    """Ends the run with one of EXIT_CODES and a message for the person reading the console."""

    def __init__(self, outcome: str, message: str):
        super().__init__(message)
        self.outcome = outcome


def progress(step: str) -> None:
    """One machine-readable line per step, e.g. `PROGRESS searching` or `PROGRESS reading 3/20`."""
    print(f"PROGRESS {step}", flush=True)


# --- the steps ------------------------------------------------------------------------------


def open_external_search(page: Page, app_url: str, headless: bool) -> None:
    """Land on AI Search -> External, logging in first if the saved session has expired."""
    search_url = f"{app_url}/resume?tab=integrated"
    page.goto(search_url)
    page.wait_for_function(  # the app decides after loading: the search page, or /login
        "location.pathname.startsWith('/login')"
        " || !!document.querySelector('#integrated-search-query')",
        timeout=PAGE_TIMEOUT * 1000,
    )
    if on_login_page(page):
        progress("logging_in")
        log_in(page, headless)
        page.goto(search_url)
    page.locator("#integrated-search-query").wait_for(timeout=PAGE_TIMEOUT * 1000)


def on_login_page(page: Page) -> bool:
    return urlparse(page.url).path.startswith("/login")


def log_in(page: Page, headless: bool) -> None:
    """Excel Hire's login is two steps: email + Login, then the password. Both are filled from
    .env. If a step does not appear as expected, the window stays open for you to finish by hand;
    either way the session is saved for the next run."""
    email = os.getenv("EXCELHIRE_EMAIL", "")
    password = os.getenv("EXCELHIRE_PASSWORD", "")
    if email and password:
        try:
            page.locator(EMAIL_INPUT).first.fill(email, timeout=10_000)
            password_input = page.locator("input[type=password]").first
            if not password_input.is_visible():
                page.locator(LOGIN_BUTTON).first.click(timeout=10_000)
                password_input.wait_for(timeout=15_000)
            password_input.fill(password)
            page.locator(LOGIN_BUTTON).first.click(timeout=10_000)
        except PlaywrightTimeout:
            print("Could not complete the login form automatically.")
    wait_for_login = 30 if headless else LOGIN_TIMEOUT
    if not headless:
        print(
            f"If you are not logged in yet, finish logging in in the browser window "
            f"({wait_for_login} s)."
        )
    try:
        page.wait_for_url(
            lambda url: not urlparse(url).path.startswith("/login"), timeout=wait_for_login * 1000
        )
    except PlaywrightTimeout:
        raise RunStopped(
            "login_failed",
            "Login did not complete. Check EXCELHIRE_EMAIL / EXCELHIRE_PASSWORD in .env, or run "
            "once with --show and log in in the window.",
        ) from None


CREDITS_PATTERN = re.compile(r"Available credits\s*(\d+)")
CREDITS_TIMEOUT = 15  # seconds; right after a login the counter can take a moment to appear


def available_credits(page: Page) -> int | None:
    """The credit counter as shown, waiting briefly for it; None if it never appears."""
    try:
        page.get_by_text(re.compile(r"Available credits")).first.wait_for(
            timeout=CREDITS_TIMEOUT * 1000
        )
    except PlaywrightTimeout:
        return None
    match = CREDITS_PATTERN.search(page.inner_text("body"))
    return int(match.group(1)) if match else None


def run_search(page: Page, jd: str, mode: str) -> str:
    """Choose the mode, paste the JD, wait for Excel Hire to turn it into filters, run the search.
    Returns the "Total matches" figure as shown."""
    progress("parsing")
    page.locator("button.search-mode__option", has_text=mode).first.click()
    page.locator("#integrated-search-query").fill(jd)

    filters_ready = page.get_by_text("Filters ready from your query")
    try:  # Excel Hire parses by itself once typing pauses; the button is the fallback
        filters_ready.wait_for(timeout=20_000)
    except PlaywrightTimeout:
        page.get_by_role("button", name="Parse Filters").click()
        filters_ready.wait_for(timeout=FILTERS_TIMEOUT * 1000)

    progress("searching")
    page.locator("button.integrated-btn--primary", has_text="Run Search").click()
    total = page.locator(".integrated-search-total-pill")
    total.wait_for(timeout=RESULTS_TIMEOUT * 1000)
    page.locator(f"{RESULT_ROW} a.candidate-link").first.wait_for(timeout=PAGE_TIMEOUT * 1000)
    return total.inner_text().replace("Total matches:", "").strip()


def read_candidates(page: Page, count: int) -> list[dict[str, str]]:
    """The first `count` rows, exactly as the table shows them once each "+N more" is expanded.
    Email and phone stay hidden behind Reveal, so they are not read."""
    rows = page.locator(RESULT_ROW)
    check_table_shape(rows.first)
    candidates = []
    wanted = min(count, rows.count())
    for index in range(wanted):
        progress(f"reading {index + 1}/{wanted}")
        row = rows.nth(index)
        name_cell = row.locator(CELL["name"])
        score_cell = row.locator(CELL["score"])
        candidates.append(
            {
                "Name": text_of(name_cell.locator("a.candidate-link")),
                "LinkedIn": name_cell.locator("a.candidate-link").get_attribute("href") or "",
                "Location": text_of(name_cell.locator(".location-text")),
                "Job Title": text_of(row.locator(CELL["job_title"])),
                "Company": text_of(row.locator(CELL["company"])),
                "Experience": "\n".join(expanded_lines(row.locator(CELL["experience"]))),
                "Skills": ", ".join(expanded_lines(row.locator(CELL["skills"]))),
                "Education": "\n".join(lines_of(row.locator(CELL["education"]))),
                "Score (%)": text_of(score_cell.locator(".score-chip")),
                "Why matched": " ".join(
                    (score_cell.locator(".match-reason-tooltip").text_content() or "").split()
                ),
            }
        )
    return candidates


def save_excel(candidates: list[dict[str, str]], jd: str, mode: str, total: str) -> Path:
    workbook = Workbook()
    write_candidates_sheet(workbook.active, candidates)
    write_search_sheet(workbook.create_sheet("Search"), jd, mode, total, len(candidates))
    OUTPUT_FOLDER.mkdir(exist_ok=True)
    path = OUTPUT_FOLDER / f"excelhire_{slug(jd)}_{datetime.now():%Y%m%d-%H%M%S}.xlsx"
    workbook.save(path)
    return path


def save_csv(candidates: list[dict[str, str]], path: Path) -> Path:
    """The LeadOps CSV: CSV_COLUMNS, one row per candidate, UTF-8. Skills stay comma-separated;
    experience, education and the match reasons keep one entry per line."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=CSV_COLUMNS)
        writer.writeheader()
        for candidate in candidates:
            writer.writerow(csv_row(candidate))
    return path


def csv_row(candidate: dict[str, str]) -> dict[str, str | int]:
    first_name, last_name = split_name(candidate["Name"])
    return {
        "first_name": first_name,
        "last_name": last_name,
        "linkedin_url": candidate["LinkedIn"],
        "location": candidate["Location"],
        "company": candidate["Company"],
        "skills": candidate["Skills"],
        "experience": candidate["Experience"],
        "education": candidate["Education"],
        "match_score": excel_value(candidate, "Score (%)", 0),
        "match_reasons": excel_value(candidate, "Why matched", 0),
    }


def split_name(name: str) -> tuple[str, str]:
    """'Y Rajesh' -> ('Y', 'Rajesh'); 'Anu' -> ('Anu', ''); extra words stay in the last name."""
    first, _, last = " ".join(name.split()).partition(" ")
    return first, last


def write_candidates_sheet(sheet: Worksheet, candidates: list[dict[str, str]]) -> None:
    sheet.title = "Candidates"
    sheet.append([header for header, _, _ in EXCEL_COLUMNS])
    for number, candidate in enumerate(candidates, start=1):
        sheet.append([excel_value(candidate, field, number) for _, _, field in EXCEL_COLUMNS])

    for column, (_, width, _) in enumerate(EXCEL_COLUMNS, start=1):
        letter = get_column_letter(column)
        sheet.column_dimensions[letter].width = width
        heading = sheet.cell(1, column)
        heading.font = Font(name=FONT, bold=True, color="FFFFFF")
        heading.fill = HEADER_FILL
        centred = "center" if width < 10 else "left"
        heading.alignment = Alignment(vertical="center", horizontal=centred)
    sheet.row_dimensions[1].height = 24

    columns = {header: index for index, (header, _, _) in enumerate(EXCEL_COLUMNS)}
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = ROW_HEIGHT
        band = BAND_FILL if row[0].row % 2 == 0 else None
        for cell in row:
            cell.font = Font(name=FONT, size=10)
            cell.alignment = TOP_LEFT
            if band:
                cell.fill = band
        row[columns["#"]].alignment = Alignment(vertical="top", horizontal="center")
        row[columns["Name"]].font = Font(name=FONT, size=10, bold=True)
        score = row[columns["Score"]]
        score.alignment = Alignment(vertical="top", horizontal="center")
        score.font = Font(name=FONT, size=11, bold=True)
        linkedin = row[columns["LinkedIn"]]
        if linkedin.value:
            linkedin.hyperlink = linkedin.value
            linkedin.value = "View profile"
            linkedin.font = Font(name=FONT, size=10, color=LINK_COLOR, underline="single")

    last_row = sheet.max_row
    if last_row > 1:  # green for the best matches, red for the weakest, on this page of results
        score_letter = get_column_letter(columns["Score"] + 1)
        sheet.conditional_formatting.add(
            f"{score_letter}2:{score_letter}{last_row}",
            ColorScaleRule(
                start_type="min",
                start_color="F8B4B4",
                mid_type="percentile",
                mid_value=50,
                mid_color="FDE68A",
                end_type="max",
                end_color="A7F3D0",
            ),
        )
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(EXCEL_COLUMNS))}{last_row}"
    sheet.freeze_panes = FROZEN_UP_TO
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def write_search_sheet(sheet: Worksheet, jd: str, mode: str, total: str, saved: int) -> None:
    for label, value in (
        ("Run at", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Search mode", mode),
        ("Total matches", total),
        ("Profiles saved", saved),
        ("Job description", jd),
    ):
        sheet.append([label, value])
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 110
    for label, value in sheet.iter_rows():
        label.font = Font(name=FONT, bold=True)
        label.alignment = Alignment(vertical="top")
        value.font = Font(name=FONT)
        value.alignment = TOP_LEFT


def excel_value(candidate: dict[str, str], field: str | None, number: int) -> str | int:
    """What one cell shows: the row number, the score as a number (so it sorts and colours), the
    "why matched" note one reason per line, everything else as read."""
    if field is None:
        return number
    value = candidate[field]
    if field == "Score (%)":
        return int(value) if value.isdigit() else value
    if field == "Why matched":
        reasons = value.split("•")  # "This candidate (…) was matched because: • ✓ Title … • ⚠ …"
        return "\n".join(reason.strip() for reason in reasons[1:] if reason.strip()) or value
    return value


# --- small helpers --------------------------------------------------------------------------


def check_table_shape(first_row: Locator) -> None:
    """Stop, rather than write the wrong columns, if Excel Hire has changed its table."""
    missing = [field for field, selector in CELL.items() if not first_row.locator(selector).count()]
    if missing or not first_row.locator("a.candidate-link").count():
        raise RunStopped(
            "table_changed",
            f"The results table has changed (missing: {', '.join(missing) or 'name link'}).",
        )


def expanded_lines(cell: Locator) -> list[str]:
    """Click the cell's own "+N more" link first: it only reveals the rest of this list in the
    page. It is a different control from Reveal, which lives in the Actions column and is never
    touched."""
    expand = cell.locator("a", has_text=EXPAND_LINK)
    if expand.count():
        expand.first.click()
        cell.locator("a", has_text=EXPAND_LINK).wait_for(state="detached", timeout=5_000)
    return [line for line in lines_of(cell) if line.lower() != "show less"]


def lines_of(cell: Locator) -> list[str]:
    lines = [line.strip() for line in cell.inner_text().splitlines()]
    return [line for line in lines if line and line != EMPTY_CELL]


def text_of(locator: Locator) -> str:
    if not locator.count():
        return ""
    text = " ".join(locator.first.inner_text().split())
    return "" if text == EMPTY_CELL else text


def slug(jd: str) -> str:
    words = re.findall(r"[A-Za-z0-9]+", jd)[:5]
    return "_".join(words).lower() or "search"


# --- entry point ----------------------------------------------------------------------------


def main() -> None:
    parser = argparse.ArgumentParser(description="Excel Hire External AI Search -> Excel")
    parser.add_argument("jd_file", type=Path, help="text file holding the job description")
    parser.add_argument("--count", type=int, default=20, help="profiles to save (max 25)")
    parser.add_argument("--mode", choices=SEARCH_MODES, default="Best Fit")
    parser.add_argument(
        "--show", action="store_true", help="show the browser window (hidden by default)"
    )
    parser.add_argument(
        "--csv", type=Path, help="also write the LeadOps CSV (CSV_COLUMNS) to this path"
    )
    args = parser.parse_args()
    try:
        outcome = search(args)
    except RunStopped as stopped:
        print(stopped, flush=True)
        outcome = stopped.outcome
    except PlaywrightTimeout as error:
        print(f"Excel Hire did not answer in time: {error}", flush=True)
        outcome = "timed_out"
    sys.exit(EXIT_CODES[outcome])


def search(args: argparse.Namespace) -> str:
    """The whole run; returns how it ended (a key of EXIT_CODES)."""
    if not args.jd_file.is_file():
        raise RunStopped("bad_input", f"No such file: {args.jd_file}")
    jd = args.jd_file.read_text(encoding="utf-8").strip()
    if len(jd) < MIN_JD_LENGTH:
        raise RunStopped(
            "bad_input", f"The job description must be at least {MIN_JD_LENGTH} characters."
        )
    if args.count > ROWS_ON_FIRST_PAGE:
        print(f"Only the first page is read: at most {ROWS_ON_FIRST_PAGE} profiles.")

    load_dotenv(HERE / ".env")
    app_url = os.getenv("EXCELHIRE_URL", "https://product.excelhire.com").rstrip("/")

    progress("opening")
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch_persistent_context(
            PROFILE_FOLDER, headless=not args.show, viewport={"width": 1600, "height": 900}
        )
        try:
            page = browser.pages[0] if browser.pages else browser.new_page()
            open_external_search(page, app_url, headless=not args.show)
            credits_before = available_credits(page)
            total = run_search(page, jd, args.mode)
            candidates = read_candidates(page, args.count)
            credits_after = available_credits(page)
        finally:
            browser.close()

    progress("saving")
    path = save_excel(candidates, jd, args.mode, total)
    print(f"{len(candidates)} profiles saved to {path}  (total matches: {total})")
    if args.csv:
        print(f"LeadOps CSV saved to {save_csv(candidates, args.csv)}")
    if credits_before is None or credits_after is None:
        # A counter that could not be read is not a change; say so instead of raising an alarm.
        print(f"Credits could not be verified (before: {credits_before}, after: {credits_after}).")
    elif credits_before != credits_after:
        print(f"WARNING: credits changed from {credits_before} to {credits_after}.")
        return "credits_changed"
    else:
        print(f"Credits unchanged: {credits_after}.")
    progress("done")
    return "ok"


if __name__ == "__main__":
    main()
